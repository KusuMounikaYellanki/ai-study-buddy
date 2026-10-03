from fastapi import FastAPI, Request, Depends, HTTPException, Header, UploadFile, File, Form
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from sse_starlette.sse import EventSourceResponse
from pydantic import BaseModel
from dotenv import load_dotenv
import httpx
import json
import os
import psycopg2
from psycopg2.extras import RealDictCursor
import jwt
import asyncio
import uuid
import io
import re
from pathlib import Path
from datetime import datetime, timedelta, timezone
import chromadb
from slowapi import Limiter
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded

load_dotenv()
app = FastAPI()
app.mount("/static", StaticFiles(directory="static"), name="static")

limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter

@app.exception_handler(RateLimitExceeded)
async def rate_limit_handler(request, exc):
    return JSONResponse(status_code=429, content={"error": "Too many requests. Please slow down."})

chroma_client = chromadb.PersistentClient(path="./chroma_data")
collection = chroma_client.get_or_create_collection(name="study_notes")

AI_PROVIDER = os.getenv("AI_PROVIDER", "ollama")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GEMINI_MODEL = "gemini-flash-latest"
GEMINI_EMBED_MODEL = "gemini-embedding-001"
SECRET_KEY = os.getenv("SECRET_KEY", "dev-only-insecure-key")
JWT_ALGORITHM = "HS256"
JWT_EXPIRY_DAYS = 30


# ---------- Database setup (Neon PostgreSQL) ----------
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
if not DATABASE_URL:
    raise RuntimeError("DATABASE_URL is not configured. Add your Neon PostgreSQL connection string to the environment.")

class DBConnection:
    """Small compatibility wrapper so the existing app can keep using conn.execute(...)."""
    def __init__(self):
        self.conn = psycopg2.connect(DATABASE_URL, cursor_factory=RealDictCursor)

    def execute(self, sql, params=None):
        # Existing queries use SQLite's '?' placeholders. PostgreSQL uses '%s'.
        sql = sql.replace("?", "%s")
        cur = self.conn.cursor()
        cur.execute(sql, params or ())
        return cur

    def commit(self):
        self.conn.commit()

    def rollback(self):
        self.conn.rollback()

    def close(self):
        self.conn.close()


def get_db():
    return DBConnection()


def init_db():
    conn = get_db()
    try:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id TEXT PRIMARY KEY,
                email TEXT UNIQUE NOT NULL,
                name TEXT,
                password_hash TEXT,
                created_at TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS chats (
                id TEXT PRIMARY KEY,
                user_id TEXT NOT NULL,
                title TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS messages (
                id TEXT PRIMARY KEY,
                chat_id TEXT NOT NULL,
                role TEXT NOT NULL,
                text TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS chat_sources (
                chat_id TEXT NOT NULL,
                source_id TEXT NOT NULL,
                source_name TEXT NOT NULL,
                attached_at TEXT NOT NULL,
                PRIMARY KEY (chat_id, source_id)
            )
        """)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


init_db()


@app.get("/")
def serve_frontend():
    return FileResponse("static/index.html")


# ---------- Auth ----------
class LoginInput(BaseModel):
    name: str
    email: str

def create_token(user_id: str) -> str:
    payload = {
        "user_id": user_id,
        "exp": datetime.now(timezone.utc) + timedelta(days=JWT_EXPIRY_DAYS),
    }
    return jwt.encode(payload, SECRET_KEY, algorithm=JWT_ALGORITHM)

def decode_token(token: str) -> str:
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[JWT_ALGORITHM])
        return payload["user_id"]
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Invalid or expired token")


async def get_current_user(
    authorization: str = Header(default=None),
    token: str = None,  # fallback for EventSource, which can't send custom headers
):
    raw_token = None
    if authorization and authorization.startswith("Bearer "):
        raw_token = authorization.split(" ", 1)[1]
    elif token:
        raw_token = token

    if not raw_token:
        raise HTTPException(status_code=401, detail="Not authenticated")

    user_id = decode_token(raw_token)

    conn = get_db()
    user = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    conn.close()
    if not user:
        raise HTTPException(status_code=401, detail="User not found")

    return user_id


@app.post("/auth/login")
def login(payload: LoginInput):
    name = payload.name.strip()
    email = payload.email.strip().lower()
    if not name or len(name) > 80:
        raise HTTPException(status_code=400, detail="Please enter your name")
    if not email or "@" not in email:
        raise HTTPException(status_code=400, detail="Please enter a valid email")

    conn = get_db()
    user = conn.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()

    if user:
        user_id = user["id"]
        conn.execute("UPDATE users SET name = ? WHERE id = ?", (name, user_id))
    else:
        user_id = str(uuid.uuid4())
        conn.execute(
            "INSERT INTO users (id, email, name, password_hash, created_at) VALUES (?, ?, ?, NULL, ?)",
            (user_id, email, name, datetime.now(timezone.utc).isoformat()),
        )
    conn.commit()
    conn.close()
    token = create_token(user_id)
    return {"token": token, "email": email, "name": name}


# ---------- Chat history endpoints ----------
def attach_pending_sources(user_id: str, chat_id: str):
    pending_tag = f"pending:{user_id}"
    data = collection.get(where={"$and": [{"session_id": user_id}, {"chat_id": pending_tag}]}, include=["metadatas"])
    source_names = {}
    for metadata in data.get("metadatas") or []:
        if metadata and metadata.get("source_id"):
            source_names[metadata["source_id"]] = metadata.get("source_name", "Study notes")

    if not source_names:
        return 0

    for source_id, source_name in source_names.items():
        source_data = collection.get(where={"$and": [{"session_id": user_id}, {"chat_id": pending_tag}, {"source_id": source_id}]}, include=["metadatas"])
        ids = source_data.get("ids") or []
        metas = source_data.get("metadatas") or []
        updated = []
        for metadata in metas:
            m = dict(metadata or {})
            m["chat_id"] = chat_id
            updated.append(m)
        if ids:
            collection.update(ids=ids, metadatas=updated)

    conn = get_db()
    now = datetime.now(timezone.utc).isoformat()
    for source_id, source_name in source_names.items():
        conn.execute("INSERT INTO chat_sources (chat_id, source_id, source_name, attached_at) VALUES (?, ?, ?, ?) ON CONFLICT (chat_id, source_id) DO NOTHING", (chat_id, source_id, source_name, now))
    conn.commit()
    conn.close()
    return len(source_names)


@app.post("/chats")
def create_chat(user_id: str = Depends(get_current_user)):
    chat_id = str(uuid.uuid4())
    conn = get_db()
    conn.execute(
        "INSERT INTO chats (id, user_id, title, created_at) VALUES (?, ?, ?, ?)",
        (chat_id, user_id, "New Chat", datetime.now(timezone.utc).isoformat()),
    )
    conn.commit()
    conn.close()
    attached = attach_pending_sources(user_id, chat_id)
    return {"chat_id": chat_id, "attached_sources": attached}


@app.get("/chats")
def list_chats(user_id: str = Depends(get_current_user)):
    conn = get_db()
    rows = conn.execute(
        "SELECT id, title, created_at FROM chats WHERE user_id = ? ORDER BY created_at DESC",
        (user_id,),
    ).fetchall()
    conn.close()
    return {"chats": [dict(r) for r in rows]}


@app.get("/chats/{chat_id}/messages")
def get_chat_messages(chat_id: str, user_id: str = Depends(get_current_user)):
    conn = get_db()
    chat = conn.execute("SELECT * FROM chats WHERE id = ? AND user_id = ?", (chat_id, user_id)).fetchone()
    if not chat:
        conn.close()
        raise HTTPException(status_code=404, detail="Chat not found")
    rows = conn.execute(
        "SELECT role, text, created_at FROM messages WHERE chat_id = ? ORDER BY created_at ASC",
        (chat_id,),
    ).fetchall()
    conn.close()
    return {"messages": [dict(r) for r in rows]}


@app.get("/chats/{chat_id}/sources")
def get_chat_sources(chat_id: str, user_id: str = Depends(get_current_user)):
    conn = get_db()
    chat = conn.execute("SELECT id FROM chats WHERE id = ? AND user_id = ?", (chat_id, user_id)).fetchone()
    if not chat:
        conn.close()
        raise HTTPException(status_code=404, detail="Chat not found")
    rows = conn.execute("SELECT source_id, source_name, attached_at FROM chat_sources WHERE chat_id = ? ORDER BY attached_at ASC", (chat_id,)).fetchall()
    conn.close()
    return {"sources": [dict(r) for r in rows]}


def save_message(chat_id: str, role: str, text: str):
    conn = get_db()
    conn.execute(
        "INSERT INTO messages (id, chat_id, role, text, created_at) VALUES (?, ?, ?, ?, ?)",
        (str(uuid.uuid4()), chat_id, role, text, datetime.now(timezone.utc).isoformat()),
    )
    # Auto-title the chat from the first user message
    if role == "user":
        count = conn.execute("SELECT COUNT(*) as c FROM messages WHERE chat_id = ? AND role = 'user'", (chat_id,)).fetchone()["c"]
        if count == 1:
            title = text[:28] + "…" if len(text) > 28 else text
            conn.execute("UPDATE chats SET title = ? WHERE id = ?", (title, chat_id))
    conn.commit()
    conn.close()


# ---------- Chunking ----------
def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks


# ---------- Embeddings (provider-aware) ----------
async def get_embedding(text: str) -> list[float]:
    try:
        async with httpx.AsyncClient(timeout=15.0) as client:
            if AI_PROVIDER == "gemini":
                response = await client.post(
                    f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_EMBED_MODEL}:embedContent?key={GEMINI_API_KEY}",
                    json={"content": {"parts": [{"text": text}]}},
                )
                response.raise_for_status()
                return response.json()["embedding"]["values"]
            else:
                response = await client.post(
                    "http://localhost:11434/api/embeddings",
                    json={"model": "nomic-embed-text", "prompt": text},
                )
                response.raise_for_status()
                return response.json()["embedding"]
    except (httpx.HTTPError, httpx.TimeoutException, KeyError) as e:
        raise RuntimeError(f"Embedding request failed: {e}")


async def retrieve_context(question: str, user_id: str, chat_id: str = None, top_k: int = 3) -> list[str]:
    query_embedding = await get_embedding(question)

    if chat_id:
        where_filter = {"$and": [{"session_id": user_id}, {"chat_id": chat_id}]}
    else:
        where_filter = {"session_id": user_id}

    results = collection.query(
        query_embeddings=[query_embedding],
        n_results=top_k,
        where=where_filter,
    )
    return results["documents"][0] if results["documents"] else []


def build_rag_prompt(question: str, context_chunks: list[str], instructions: str = None) -> str:
    context_text = "\n\n".join(context_chunks) if context_chunks else "No notes available."
    task = instructions or "Answer the student's question using the notes. If the notes don't cover it, say so."
    return f"""You are a helpful study assistant.

Notes:
{context_text}

Task: {task}

Question/Topic: {question}

Important: Give ONLY the final answer. Do not think out loud or narrate your reasoning process. Be direct and concise.

Response:"""


# ---------- Streaming generation (provider-aware, now saves to DB) ----------
async def ai_token_stream(prompt: str, chat_id: str = None):
    full_response = ""
    try:
        if AI_PROVIDER == "gemini":
            async with httpx.AsyncClient(timeout=30.0) as client:
                for attempt in range(3):
                    async with client.stream(
                        "POST",
                        f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:streamGenerateContent?alt=sse&key={GEMINI_API_KEY}",
                        json={"contents": [{"parts": [{"text": prompt}]}]},
                    ) as response:
                        if response.status_code in (503, 429) and attempt < 2:
                            print(f"[retry] Gemini attempt {attempt + 1} got {response.status_code}, retrying", flush=True)
                            await asyncio.sleep(2 ** attempt)
                            continue
                        response.raise_for_status()
                        async for line in response.aiter_lines():
                            if line.startswith("data: "):
                                payload = line[len("data: "):].strip()
                                if not payload or payload == "[DONE]":
                                    continue
                                try:
                                    chunk = json.loads(payload)
                                    text = chunk["candidates"][0]["content"]["parts"][0]["text"]
                                    full_response += text
                                    yield {"data": text}
                                except (KeyError, IndexError, json.JSONDecodeError):
                                    continue
                        break
        else:
            async with httpx.AsyncClient(timeout=30.0) as client:
                async with client.stream(
                    "POST",
                    "http://localhost:11434/api/generate",
                    json={"model": "llama3.1:8b", "prompt": prompt, "stream": True},
                ) as response:
                    response.raise_for_status()
                    async for line in response.aiter_lines():
                        if line:
                            chunk = json.loads(line)
                            text = chunk.get("response", "")
                            full_response += text
                            yield {"data": text}
    except (httpx.HTTPError, httpx.TimeoutException) as e:
        print(f"[ai_token_stream ERROR - HTTP/Timeout] {e}", flush=True)
        full_response = "⚠️ Sorry, the AI service is currently unavailable. Please try again in a moment."
        yield {"data": full_response}
    except Exception as e:
        print(f"[ai_token_stream ERROR - Unexpected] {type(e).__name__}: {e}", flush=True)
        full_response = "⚠️ Something went wrong while generating a response. Please try again."
        yield {"data": full_response}
    finally:
        if chat_id and full_response:
            save_message(chat_id, "ai", full_response)


async def no_notes_stream(chat_id: str = None):
    msg = "You haven't uploaded any notes yet! Please paste some study notes above and click 'Upload Notes' first, then try again."
    yield {"data": msg}
    if chat_id:
        save_message(chat_id, "ai", msg)


# ---------- Tool definitions ----------
TOOLS_OLLAMA = [
    {"type": "function", "function": {
        "name": "generate_flashcards",
        "description": "Generate study flashcards (question/answer pairs) about a topic",
        "parameters": {"type": "object", "properties": {"topic": {"type": "string"}}, "required": ["topic"]},
    }},
    {"type": "function", "function": {
        "name": "generate_quiz",
        "description": "Generate multiple-choice quiz questions about a topic",
        "parameters": {"type": "object", "properties": {
            "topic": {"type": "string"}, "num_questions": {"type": "integer"}}, "required": ["topic"]},
    }},
    {"type": "function", "function": {
        "name": "summarize_notes",
        "description": "Summarize the stored notes on a topic",
        "parameters": {"type": "object", "properties": {"topic": {"type": "string"}}, "required": ["topic"]},
    }},
]

GEMINI_TOOLS = [{"functionDeclarations": [
    {"name": "generate_flashcards", "description": "Generate study flashcards (question/answer pairs) about a topic",
     "parameters": {"type": "object", "properties": {"topic": {"type": "string"}}, "required": ["topic"]}},
    {"name": "generate_quiz", "description": "Generate multiple-choice quiz questions about a topic",
     "parameters": {"type": "object", "properties": {
         "topic": {"type": "string"}, "num_questions": {"type": "integer"}}, "required": ["topic"]}},
    {"name": "summarize_notes", "description": "Summarize the stored notes on a topic",
     "parameters": {"type": "object", "properties": {"topic": {"type": "string"}}, "required": ["topic"]}},
]}]

ROUTING_INSTRUCTIONS = (
    "You are a routing assistant. Only call a tool/function if the user's message contains an "
    "explicit request word: 'flashcards', 'quiz', 'summarize', or 'summary'. If none of those "
    "exact words appear in the message, do NOT call any tool, no matter how much content exists "
    "on the topic.\n\n"
    "Examples:\n"
    "'What is seed?' -> NO tool call, answer directly.\n"
    "'What is photosynthesis?' -> NO tool call, answer directly.\n"
    "'Summarize seed' -> call summarize_notes.\n"
    "'Give me flashcards on seed' -> call generate_flashcards.\n"
    "'Quiz me on seed' -> call generate_quiz."
)


async def generate_structured_action(action: str, topic: str, context_chunks: list[str]):
    """Generate structured data for interactive flashcards or quizzes."""
    context_text = "\n\n".join(context_chunks) if context_chunks else "No notes available."
    if action == "flashcards":
        schema_hint = '{"cards":[{"question":"...","answer":"..."}]}'
        task = (
            "Create exactly 5 study flashcards using ONLY the supplied notes. "
            "Each card must have a concise question and an accurate answer supported by the notes. "
            f"Return ONLY valid JSON matching this shape: {schema_hint}"
        )
    else:
        schema_hint = '{"questions":[{"question":"...","options":["...","...","...","..."],"answer":0,"explanation":"..."}]}'
        task = (
            "Create exactly 5 multiple-choice quiz questions using ONLY the supplied notes. "
            "Each question must have exactly 4 options. 'answer' is the zero-based index of the correct option. "
            "Include a short explanation supported by the notes. "
            f"Return ONLY valid JSON matching this shape: {schema_hint}"
        )
    prompt = f"""You are a study-material generator.\n\nStudy notes:\n{context_text}\n\nRequested notes: {topic}\n\nTask: {task}\nDo not add facts that are not supported by the notes."""
    try:
        async with httpx.AsyncClient(timeout=60.0) as client:
            if AI_PROVIDER == "gemini":
                response = await client.post(
                    f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent?key={GEMINI_API_KEY}",
                    json={
                        "contents": [{"parts": [{"text": prompt}]}],
                        "generationConfig": {"responseMimeType": "application/json"},
                    },
                )
                response.raise_for_status()
                data = response.json()
                raw = data["candidates"][0]["content"]["parts"][0]["text"]
            else:
                response = await client.post(
                    "http://localhost:11434/api/generate",
                    json={"model": "llama3.1:8b", "prompt": prompt, "format": "json", "stream": False},
                )
                response.raise_for_status()
                raw = response.json().get("response", "")
        raw = raw.strip()
        if raw.startswith("```"):
            raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE | re.DOTALL).strip()
        return json.loads(raw)
    except Exception as e:
        print(f"[structured action ERROR] {type(e).__name__}: {e}", flush=True)
        return None


async def structured_action_stream(action: str, topic: str, context_chunks: list[str], chat_id: str):
    payload = await generate_structured_action(action, topic, context_chunks)
    if not payload:
        message = "⚠️ I couldn't create that study activity right now. Please try again."
        save_message(chat_id, "ai", message)
        yield {"data": message}
        return

    if action == "flashcards":
        cards = payload.get("cards", [])
        cards = cards[:5]
        text = "Here are your flashcards:\n\n" + "\n\n".join(
            f"{i+1}. Q: {c.get('question','')}\nA: {c.get('answer','')}" for i, c in enumerate(cards)
        )
    else:
        questions = payload.get("questions", [])[:5]
        text = "Here is your quiz:\n\n" + "\n\n".join(
            f"{i+1}. {q.get('question','')}\n" + "\n".join(f"{chr(65+j)}. {opt}" for j, opt in enumerate(q.get('options', [])[:4]))
            for i, q in enumerate(questions)
        )
    save_message(chat_id, "ai", text)
    event = {"__study_action__": action, "topic": topic, "data": payload}
    yield {"data": json.dumps(event)}


async def decide_tool(message: str):
    if AI_PROVIDER == "gemini":
        async with httpx.AsyncClient(timeout=None) as client:
            response = await client.post(
                f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent?key={GEMINI_API_KEY}",
                json={
                    "contents": [{"parts": [{"text": message}]}],
                    "tools": GEMINI_TOOLS,
                    "systemInstruction": {"parts": [{"text": ROUTING_INSTRUCTIONS}]},
                },
            )
            data = response.json()
            try:
                for part in data["candidates"][0]["content"]["parts"]:
                    if "functionCall" in part:
                        fc = part["functionCall"]
                        return [{"function": {"name": fc["name"], "arguments": fc.get("args", {})}}]
            except (KeyError, IndexError):
                pass
            return None
    else:
        async with httpx.AsyncClient(timeout=None) as client:
            response = await client.post(
                "http://localhost:11434/api/chat",
                json={
                    "model": "llama3.1:8b",
                    "messages": [
                        {"role": "system", "content": ROUTING_INSTRUCTIONS},
                        {"role": "user", "content": message},
                    ],
                    "tools": TOOLS_OLLAMA,
                    "stream": False,
                },
            )
            data = response.json()
            return data.get("message", {}).get("tool_calls")


# ---------- File/text extraction ----------
MAX_UPLOAD_SIZE = 10 * 1024 * 1024  # 10 MB
ALLOWED_EXTENSIONS = {".pdf", ".txt", ".doc", ".docx", ".xls", ".xlsx"}


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    """Extract readable text from supported study-file formats."""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Use PDF, TXT, DOC, DOCX, XLS, or XLSX.")

    if ext == ".txt":
        return data.decode("utf-8", errors="replace")

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages)

    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    parts.append(" | ".join(values))
        return "\n".join(parts)

    if ext == ".xls":
        raise ValueError("Legacy .xls files are not supported yet. Please save the file as .xlsx and upload it again.")

    if ext == ".doc":
        raise ValueError("Legacy .doc files are not supported directly. Please save the file as .docx and upload it again.")

    raise ValueError("Unsupported file type")


async def store_notes_text(
    text: str,
    user_id: str,
    chat_id: str | None = None,
    source_name: str = "Pasted notes",
) -> int:
    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in the supplied content.")

    chunks = chunk_text(text)
    chat_tag = chat_id or f"pending:{user_id}"
    source_id = str(uuid.uuid4())
    uploaded_at = datetime.now(timezone.utc).isoformat()
    ids, embeddings, metadatas = [], [], []
    for chunk_index, chunk in enumerate(chunks):
        embedding = await get_embedding(chunk)
        ids.append(str(uuid.uuid4()))
        embeddings.append(embedding)
        metadatas.append({
            "session_id": user_id,
            "chat_id": chat_tag,
            "source_id": source_id,
            "source_name": source_name[:200],
            "uploaded_at": uploaded_at,
            "chunk_index": chunk_index,
        })

    collection.add(ids=ids, embeddings=embeddings, documents=chunks, metadatas=metadatas)
    return len(chunks)


def get_latest_source(user_id: str, chat_id: str | None = None):
    if chat_id:
        where_filter = {"$and": [{"session_id": user_id}, {"chat_id": chat_id}]}
    else:
        where_filter = {"session_id": user_id}

    data = collection.get(where=where_filter, include=["metadatas"])
    latest = None
    for metadata in data.get("metadatas") or []:
        if not metadata or not metadata.get("source_id"):
            continue
        uploaded_at = metadata.get("uploaded_at", "")
        if latest is None or uploaded_at > latest.get("uploaded_at", ""):
            latest = metadata
    return latest


def get_source_context(user_id: str, source_id: str, chat_id: str | None = None) -> list[str]:
    if chat_id:
        where_filter = {"$and": [
            {"session_id": user_id},
            {"chat_id": chat_id},
            {"source_id": source_id},
        ]}
    else:
        where_filter = {"$and": [
            {"session_id": user_id},
            {"source_id": source_id},
        ]}
    data = collection.get(where=where_filter, include=["documents", "metadatas"])
    pairs = list(zip(data.get("documents") or [], data.get("metadatas") or []))
    pairs.sort(key=lambda pair: pair[1].get("chunk_index", 0) if pair[1] else 0)
    return [doc for doc, _ in pairs]


# ---------- Notes ----------
class NotesInput(BaseModel):
    text: str
    chat_id: str | None = None
    source_name: str = "Pasted notes"


@app.post("/upload-notes")
@limiter.limit("10/minute")
async def upload_notes(request: Request, payload: NotesInput, user_id: str = Depends(get_current_user)):
    try:
        source_name = payload.source_name.strip()[:120] or "Pasted notes"
        chunks_stored = await store_notes_text(payload.text, user_id, payload.chat_id, source_name)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except RuntimeError:
        return {"error": "Failed to process notes. Please try again in a moment."}

    latest = get_latest_source(user_id, payload.chat_id)
    return {
        "message": "Notes embedded and stored",
        "chunks_stored": chunks_stored,
        "source": latest,
    }


@app.get("/notes/latest")
async def latest_notes(user_id: str = Depends(get_current_user)):
    latest = get_latest_source(user_id)
    return {"source": latest}


@app.post("/upload-file")
@limiter.limit("10/minute")
async def upload_file(
    request: Request,
    file: UploadFile = File(...),
    chat_id: str | None = Form(default=None),
    user_id: str = Depends(get_current_user),
):
    filename = file.filename or "study-file"
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Unsupported file type. Upload PDF, TXT, DOCX, or XLSX files.")

    data = await file.read()
    if len(data) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=413, detail="File is too large. Maximum size is 10 MB.")

    try:
        extracted_text = extract_text_from_bytes(filename, data)
        chunks_stored = await store_notes_text(extracted_text, user_id, chat_id, filename)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except RuntimeError:
        raise HTTPException(status_code=502, detail="Failed to process the file. Please try again in a moment.")
    except Exception as e:
        print(f"[upload-file ERROR] {type(e).__name__}: {e}", flush=True)
        raise HTTPException(status_code=400, detail="Could not read this file. Please check that it is a valid document.")

    latest = get_latest_source(user_id, chat_id)
    return {
        "message": "File uploaded and processed",
        "filename": filename,
        "chunks_stored": chunks_stored,
        "source": latest,
    }


@app.delete("/clear-notes")
def clear_notes(user_id: str = Depends(get_current_user)):
    collection.delete(where={"session_id": user_id})
    return {"message": "Your notes have been cleared"}


# ---------- Chat ----------
def get_chat_sources_list(user_id: str, chat_id: str):
    conn = get_db()
    rows = conn.execute("""
        SELECT cs.source_id, cs.source_name, cs.attached_at
        FROM chat_sources cs
        JOIN chats c ON c.id = cs.chat_id
        WHERE cs.chat_id = ? AND c.user_id = ?
        ORDER BY cs.attached_at ASC
    """, (chat_id, user_id)).fetchall()
    conn.close()
    sources = [dict(r) for r in rows]
    if sources:
        return sources

    # Backfill source associations for chats created before chat_sources existed.
    data = collection.get(where={"$and": [{"session_id": user_id}, {"chat_id": chat_id}]}, include=["metadatas"])
    discovered = {}
    for metadata in data.get("metadatas") or []:
        if metadata and metadata.get("source_id"):
            discovered[metadata["source_id"]] = {
                "source_id": metadata["source_id"],
                "source_name": metadata.get("source_name", "Study notes"),
                "attached_at": metadata.get("uploaded_at", datetime.now(timezone.utc).isoformat()),
            }
    if discovered:
        conn = get_db()
        for source in discovered.values():
            conn.execute("INSERT INTO chat_sources (chat_id, source_id, source_name, attached_at) VALUES (?, ?, ?, ?) ON CONFLICT (chat_id, source_id) DO NOTHING", (chat_id, source["source_id"], source["source_name"], source["attached_at"]))
        conn.commit()
        conn.close()
    return list(discovered.values())


def find_source_for_message(message: str, sources: list[dict]):
    for source in sources:
        name = source.get("source_name", "").lower()
        if name and name in message:
            return source
        stem = Path(name).stem.lower() if name else ""
        if stem and len(stem) > 2 and stem in message:
            return source
    return None


async def single_message_stream(message: str):
    yield {"data": message}


@app.get("/chat")
@limiter.limit("15/minute")
async def chat(request: Request, message: str, chat_id: str, user_id: str = Depends(get_current_user)):
    conn = get_db()
    chat = conn.execute("SELECT id FROM chats WHERE id = ? AND user_id = ?", (chat_id, user_id)).fetchone()
    conn.close()
    if not chat:
        raise HTTPException(status_code=404, detail="Chat not found")

    save_message(chat_id, "user", message)

    sources = get_chat_sources_list(user_id, chat_id)
    if not sources:
        return EventSourceResponse(no_notes_stream(chat_id))

    normalized = re.sub(r"\s+", " ", message.lower()).strip()
    generic_action = any(word in normalized for word in ["summarize", "summary", "quiz", "flashcards", "flash card"])
    explicit_source = find_source_for_message(normalized, sources)

    if generic_action and not explicit_source and len(sources) > 1:
        names = "\n".join(f"• {s['source_name']}" for s in sources)
        if "flashcard" in normalized:
            action_help = "Please tell me which notes you want flashcards from, for example: 'Create flashcards from Flower Study Notes'."
        elif "quiz" in normalized:
            action_help = "Please tell me which notes you want the quiz from, for example: 'Quiz me on Flower Study Notes'."
        else:
            action_help = "Please tell me which notes you want summarized, for example: 'Summarize Flower Study Notes'."
        prompt_message = f"You have {len(sources)} study materials in this chat. Which notes would you like me to use?\n\n{names}\n\n{action_help}"
        save_message(chat_id, "ai", prompt_message)
        return EventSourceResponse(single_message_stream(prompt_message))

    # Handle study actions deterministically so explicit note names always work.
    action = None
    if "flashcard" in normalized:
        action = "flashcards"
    elif re.search(r"\bquiz\b", normalized):
        action = "quiz"
    elif "summarize" in normalized or "summary" in normalized:
        action = "summarize"

    if action:
        if explicit_source:
            context_chunks = get_source_context(user_id, explicit_source["source_id"], chat_id)
            topic = explicit_source["source_name"]
        elif len(sources) == 1:
            context_chunks = get_source_context(user_id, sources[0]["source_id"], chat_id)
            topic = sources[0]["source_name"]
        else:
            names = "\n".join(f"• {s['source_name']}" for s in sources)
            if action == "flashcards":
                action_help = "Choose the notes you want flashcards from."
            elif action == "quiz":
                action_help = "Choose the notes you want the quiz from."
            else:
                action_help = "Choose the notes you want summarized."
            prompt_message = f"You have {len(sources)} study materials in this chat. Which notes would you like me to use?\n\n{names}\n\n{action_help}"
            save_message(chat_id, "ai", prompt_message)
            return EventSourceResponse(single_message_stream(prompt_message))

        if action == "flashcards":
            return EventSourceResponse(structured_action_stream("flashcards", topic, context_chunks, chat_id))
        if action == "quiz":
            return EventSourceResponse(structured_action_stream("quiz", topic, context_chunks, chat_id))

        instructions = (
            "Summarize only the selected notes. Present the summary as clear bullet points, "
            "grouped under short headings when useful. Do not write one long paragraph. "
            "Do not add information that is not supported by the selected notes."
        )
        prompt = build_rag_prompt(topic, context_chunks, instructions)
        return EventSourceResponse(ai_token_stream(prompt, chat_id))

    context_chunks = await retrieve_context(message, user_id, chat_id)
    prompt = build_rag_prompt(message, context_chunks)
    return EventSourceResponse(ai_token_stream(prompt, chat_id))
